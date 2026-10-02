"""
Deterministic Extractor Module for Aegis Knowledge Ingestion System.
Handles 100% rule-based extraction for structured and digital text files:
- JSON (configuration dumps)
- XLSX (component registers, revision timelines)
- HTML (legacy documentation)
- DOCX (glossaries, field notes)
- PPTX (operator training slides)
- Born-Digital PDFs (manuals, engineering change notices)
"""

import json
import os
import openpyxl
from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
import pypdf


class DeterministicExtractor:
    def __init__(self, base_dir: str):
        self.base_dir = base_dir

    def extract_config(self, rel_path: str = "configuration/configuration_export.json") -> dict:
        full_path = os.path.join(self.base_dir, rel_path)
        with open(full_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {
            "source_file": rel_path,
            "type": "machine_configuration",
            "content": data
        }

    def extract_xlsx(self, rel_path: str) -> dict:
        full_path = os.path.join(self.base_dir, rel_path)
        wb = openpyxl.load_workbook(full_path, data_only=True)
        sheets_data = {}
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if any(row):  # filter out completely empty rows
                    rows.append([str(c).strip() if c is not None else "" for c in row])
            if rows:
                headers = rows[0]
                records = []
                for r in rows[1:]:
                    record = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                    records.append(record)
                sheets_data[sheet_name] = {
                    "headers": headers,
                    "rows": records
                }
        return {
            "source_file": rel_path,
            "type": "spreadsheet",
            "sheets": sheets_data
        }

    def extract_html(self, rel_path: str = "manuals/legacy_manual_v1.html") -> dict:
        full_path = os.path.join(self.base_dir, rel_path)
        with open(full_path, "r", encoding="utf-8") as f:
            soup = BeautifulSoup(f.read(), "html.parser")

        title = soup.title.string if soup.title else ""
        sections = []
        current_section = {"title": "Introduction", "paragraphs": []}

        for element in soup.find_all(["h1", "h2", "h3", "p", "li"]):
            tag = element.name
            text = element.get_text(strip=True)
            if not text:
                continue
            if tag in ["h1", "h2", "h3"]:
                if current_section["paragraphs"]:
                    sections.append(current_section)
                current_section = {"title": text, "paragraphs": []}
            else:
                current_section["paragraphs"].append(text)

        if current_section["paragraphs"]:
            sections.append(current_section)

        return {
            "source_file": rel_path,
            "type": "legacy_manual_html",
            "title": title,
            "sections": sections
        }

    def extract_docx(self, rel_path: str) -> dict:
        full_path = os.path.join(self.base_dir, rel_path)
        doc = Document(full_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables_data = []
        for table in doc.tables:
            t_rows = []
            for row in table.rows:
                t_rows.append([cell.text.strip() for cell in row.cells])
            if t_rows:
                tables_data.append(t_rows)

        return {
            "source_file": rel_path,
            "type": "word_document",
            "paragraphs": paragraphs,
            "tables": tables_data
        }

    def extract_pptx(self, rel_path: str = "extra/training_slide_excerpt.pptx") -> dict:
        full_path = os.path.join(self.base_dir, rel_path)
        prs = Presentation(full_path)
        slides_data = []
        for idx, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)
            slides_data.append({
                "slide_number": idx + 1,
                "text": " ".join(slide_texts),
                "items": slide_texts
            })

        return {
            "source_file": rel_path,
            "type": "presentation_slides",
            "slides": slides_data
        }

    def extract_born_digital_pdf(self, rel_path: str) -> dict:
        full_path = os.path.join(self.base_dir, rel_path)
        reader = pypdf.PdfReader(full_path)
        pages_data = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages_data.append({
                "page_number": idx + 1,
                "text": text.strip()
            })

        return {
            "source_file": rel_path,
            "type": "born_digital_pdf",
            "total_pages": len(pages_data),
            "pages": pages_data
        }

    def is_noise_file(self, rel_path: str) -> bool:
        """Deterministic noise filter for irrelevant files (e.g. MSDS)."""
        lower = rel_path.lower()
        if "safety_data_sheet" in lower or "noise/" in lower or "msds" in lower:
            return True
        return False
